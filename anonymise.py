"""
Anonymise site names in the thesis results.

Every site is replaced by its code from SITES below ("Site A" ...): the
address in any spelling (Långebergavägen 190, LANGEBERGAVAGEN_190, ...),
the meter number in front of it, and the company name. Sites in the data
that are NOT in the list get "Other site 1", "Other site 2", ... (sorted by
name, so the same input always gives the same codes).

The originals are never changed: everything is written to a new folder.

Two modes
---------
1. Results (fast): anonymise finished output files.
       python anonymise.py results  SOURCE  DEST  [--sites path/to/site_comparison.csv]
   SOURCE is a folder (all sub-folders too) or one file. Text files (.csv,
   .json, .md, .txt, .log, .tsv, .html, .py) and Excel files (.xlsx) are
   rewritten; file and folder names are renamed. PICTURES (.png, .jpg,
   .pdf) cannot be edited -- they are NOT copied and are listed at the
   end. Use mode 2 to get anonymised figures.

2. Workspace (complete, figures included): an anonymised copy of the
   working folder -- meter files, the Solar folder, site_comparison.csv
   and the scripts (their site ids in SELECTION / CONFIG are replaced
   too). Rerun the scripts THERE and every table, plot and file name comes
   out anonymised:
       python anonymise.py workspace  SOURCE  DEST
       cd DEST
       python lpc_10.py --rerun-all     (or --rank-only)
       python filter_sensitivity.py
       python portfolio.py --mode 15min   ...

Both modes end with a LEFTOVER CHECK: every file written is searched for the
street names of the list, the company names and long meter numbers. Anything
found is printed (file, line). Nothing should be found before the results
leave your computer.

The key (real name -> code) is written to anonymisation_key.csv NEXT TO
the source, never into DEST. Keep it private.
"""
import os
import re
import shutil
import sys
import unicodedata

# ======================================================================
# THE LIST: address, company, code
# ======================================================================
CODE_PREFIX = "Site"          # "Site A", "Site B", ...
OTHER_PREFIX = "Other site"   # sites in the data that are not in the list

SITES = [
    ("Långeberga 190",   "Långeberga Förvaltning KB",           "A"),
    ("Knut Påls väg 19", "3N produkter",                        "B"),
    ("Långeberga 52",    "Klimat Transport & Logistik AB",      "C"),
    ("Långeberga 181",   "Greenfood AB",                        "D"),
    ("Långeberga 100",   "ICA DE AB",                           "E"),
    ("Kardangatan 21",   "Frigoscandia Åkeri",                  "F"),
    ("Kardangatan 11",   "BREEVA II Kamaxeln 3 AB",             "G"),
    ("Kardangatan 2",    "Castellum Långeberga AB",             "H"),
    ("Knut Påls väg 2",  "SLP Aggregatet 2 AB",                 "I"),
    ("Knut Påls väg 5",  "SLP Ackumulatorn 1 AB",               "J"),
    ("Knut Påls väg 6",  "DAD Direct/Ravn Distribution",        "K"),
    ("Knut Påls väg 7",  "Sardonyx",                            "L"),
    ("Knut Påls väg 8",  "Mixtura",                             "M"),
    ("Knut Påls väg 9",  "Sardonyx 10 AB",                      "N"),
    ("Knut Påls väg 11", "Sydgrönt Fastighets AB",              "O"),
    ("Knut Påls väg 13", "Svenska Odlarlaget Ek förening",      "P"),
    ("Knut Påls väg 15", "FG Centralen AB",                     "Q"),
    ("Dynamogatan 7",    "Castellum Långeberga AB",             "R"),
]
# A company name shared by several sites (Castellum Långeberga AB: H and R)
# cannot point to one site; it is replaced by this text instead.
SHARED_COMPANY_TEXT = "[company]"

# Where the list of sites in the data is read from (both modes): every
# site_id in these files gets a code (list sites by address, the rest
# "Other site n").
SITE_SOURCES = ["site_comparison.csv"]
# Folders whose data-file names are sites too (meter files, Solar files).
DATA_EXT = (".csv", ".xlsx", ".xls", ".xlsm", ".txt")

TEXT_EXT = (".csv", ".tsv", ".json", ".md", ".txt", ".log", ".html", ".htm", ".py", ".tex", ".bib", ".yaml", ".yml")
EXCEL_EXT = (".xlsx", ".xlsm")
PICTURE_EXT = (".png", ".jpg", ".jpeg", ".pdf", ".svg", ".gif", ".tif", ".tiff")
SKIP_DIRS = {"__pycache__", ".git", ".venv", "venv", ".idea", ".vscode"}
# Words that describe a FILE, not a site (kept, not part of the name)
FILE_WORDS = ["förbrukning", "forbrukning", "produktion", "production", "consumption", "export",
              "import", "inmatning", "uttag", "solel", "solproduktion", "pv"]
# A meter number standing alone (not in front of an address) is replaced
# only if it has at least this many digits.
MIN_BARE_NUMBER = 6
# Words dropped when they follow a code, so labels read "Site F", not
# "Site F förbrukning" (production files keep their word).
DROP_AFTER_CODE = ["förbrukning", "forbrukning", "Förbrukning", "FÖRBRUKNING", "consumption"]
# Long digit runs (meter numbers) reported by the leftover check
LEFTOVER_DIGITS = 10


# ======================================================================
# matching
# ======================================================================
_ACC = {"a": "[aåäAÅÄ]", "o": "[oöOÖ]", "e": "[eéEÉ]"}


def plain(s):
    """lower case, å/ä/ö -> a/a/o, punctuation -> space"""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"[^0-9a-z]+", " ", s.lower()).strip()


def _loose(word):
    """regex for one word that tolerates å/ä/ö written as a/o and any case"""
    out = []
    for ch in plain(word):
        out.append(_ACC.get(ch, re.escape(ch) if not ch.isalpha() else f"[{ch}{ch.upper()}]"))
    return "".join(out)


SEP = r"[\s_\-.,]*"


def address_regex(address):
    """'Långeberga 190' also matches 'Långebergavägen 190', 'LANGEBERGAVAGEN_190',
    '7359991240001455955 langebergavagen 190' (the meter number is part of
    the match); never 'Långeberga 1900' or '...19'."""
    words = plain(address).split()
    num = words[-1]
    street = words[:-1]
    pat = SEP.join(_loose(w) for w in street) + r"[A-Za-zÅÄÖåäöéÉ]*" + SEP + re.escape(num) + r"(?![0-9])"
    return re.compile(r"(?<![0-9A-Za-zÅÄÖåäö])(?:(\d{3,})" + SEP + r")?" + pat)


def company_regex(name):
    words = re.split(r"\s+", name.strip())
    pat = r"\s*".join(_loose(w) if re.fullmatch(r"[\wÅÄÖåäöÉé]+", w) else re.escape(w) for w in words)
    return re.compile(r"(?<![0-9A-Za-zÅÄÖåäö])" + pat + r"(?![0-9A-Za-zÅÄÖåäö])")


def code_of(letter):
    return f"{CODE_PREFIX} {letter}"


class Anonymiser:
    def __init__(self):
        self.addr = [(address_regex(a), code_of(c), a) for a, _, c in SITES]
        names = {}
        for _, comp, c in SITES:
            names.setdefault(comp, []).append(code_of(c))
        self.comp = [(company_regex(n), v[0] if len(v) == 1 else SHARED_COMPANY_TEXT, n)
                     for n, v in sorted(names.items(), key=lambda x: -len(x[0]))]
        self.numbers = {}       # meter number -> code
        self.other = {}         # exact name of an unlisted site -> code
        self.key_rows = []      # (real, code, how)

    # ---- learn the sites present in the data --------------------------
    def code_for_name(self, name):
        """code of a site id / file name, or None if not in the list"""
        for rx, code, _ in self.addr:
            if rx.search(name):
                return code
        return None

    def learn(self, names):
        """names: site ids and data-file names (without extension)."""
        unlisted = []
        for n in sorted(set(names)):
            n = n.strip()
            if not n:
                continue
            code = self.code_for_name(n)
            m = re.match(r"\s*(\d{3,})", n)
            if code:
                if m:
                    self.numbers.setdefault(m.group(1), code)
                self.key_rows.append((n, code, "address in the list"))
            else:
                unlisted.append(n)
        # unlisted sites: group spellings of the same site by their name
        # without number and file words (consumption + production files)
        groups = {}
        for n in unlisted:
            k = plain(self._core(n))
            if k:
                groups.setdefault(k, []).append(n)
        for i, k in enumerate(sorted(groups), 1):
            code = f"{OTHER_PREFIX} {i}"
            for n in groups[k]:
                self.other[n] = code
                core = self._core(n)
                if len(core) > 3:
                    self.other.setdefault(core, code)
                m = re.match(r"\s*(\d{3,})", n)
                if m:
                    self.numbers.setdefault(m.group(1), code)
                self.key_rows.append((n, code, "not in the list"))

    def _other_rx(self):
        if getattr(self, "_orx_n", None) != len(self.other):
            cores = {}
            for real, code in self.other.items():
                c = plain(self._core(real))
                if len(c) > 3:
                    cores[c] = code
            self._orx = [(re.compile(r"(?<![0-9A-Za-zÅÄÖåäö])(?:\d{3,}" + SEP + r")?"
                                     + SEP.join(_loose(w) for w in c.split()) + r"(?![0-9A-Za-zÅÄÖåäö])"), code)
                         for c, code in sorted(cores.items(), key=lambda x: -len(x[0]))]
            self._orx_n = len(self.other)
        return self._orx

    @staticmethod
    def _core(n):
        s = re.sub(r"^\s*\d{3,}[\s_\-]*", "", n)
        words = [w for w in re.split(r"[\s_]+", s) if plain(w) not in {plain(x) for x in FILE_WORDS}]
        return " ".join(words).strip()

    # ---- replace -------------------------------------------------------
    def text(self, s):
        for rx, code, _ in self.addr:
            s = rx.sub(code, s)
        # unlisted sites: number + name in any case / å-ä-ö spelling, longest first
        for rx, code in self._other_rx():
            s = rx.sub(code, s)
        for rx, code, _ in self.comp:
            s = rx.sub(code, s)
        # a meter number on its own (only long ones: a short number could be a kW value)
        for num in sorted(self.numbers, key=len, reverse=True):
            if len(num) >= MIN_BARE_NUMBER:
                s = re.sub(r"(?<![\d.,])" + re.escape(num) + r"(?![\d]|[.,]\d)", self.numbers[num], s)
        s = re.sub(r"((?:%s|%s) [A-Z0-9]+)[\s_]+(?:%s)(?![A-Za-zåäö])" % (
            re.escape(CODE_PREFIX), re.escape(OTHER_PREFIX), "|".join(map(re.escape, DROP_AFTER_CODE))), r"\1", s)
        # "Site A Site A" can appear when a meter number stood next to its own address
        s = re.sub(r"\b((?:%s|%s) [A-Z0-9]+)[\s_]+\1\b" % (re.escape(CODE_PREFIX), re.escape(OTHER_PREFIX)), r"\1", s)
        return s

    def leftovers(self, s):
        hits = []
        for _, _, a in self.addr:
            street = " ".join(plain(a).split()[:-1])
            rx = re.compile(SEP.join(_loose(w) for w in street.split()), re.I)
            for m in rx.finditer(s):
                hits.append(m.group(0))
        for rx, _, n in self.comp:
            hits += [m.group(0) for m in rx.finditer(s)]
        hits += re.findall(r"(?<![\d.,])\d{%d,}(?![\d]|[.,]\d)" % LEFTOVER_DIGITS, s)
        for rx, _ in self._other_rx():
            hits += [m.group(0) for m in rx.finditer(s)]
        return hits


# ======================================================================
# files
# ======================================================================
def collect_names(src):
    """site ids and data-file names found under src"""
    import csv
    names = []
    if os.path.isfile(src) and os.path.basename(src) not in SITE_SOURCES and src.lower().endswith(".csv"):
        SITE_SOURCES.append(os.path.basename(src))
    root = src if os.path.isdir(src) else os.path.dirname(os.path.abspath(src))
    for dirpath, dirs, files in os.walk(root):
        dirs[:] = [d for d in dirs if d not in SKIP_DIRS]
        for f in files:
            stem, ext = os.path.splitext(f)
            if f in SITE_SOURCES:
                p = os.path.join(dirpath, f)
                with open(p, encoding="utf-8-sig", newline="") as fh:
                    first = fh.readline()
                    sep = ";" if first.count(";") >= first.count(",") else ","
                    fh.seek(0)
                    for row in csv.DictReader(fh, delimiter=sep):
                        for col in ("site_id", "meter_file"):
                            v = row.get(col)
                            if v:
                                names.append(os.path.splitext(v)[0] if col == "meter_file" else v)
            elif ext.lower() in DATA_EXT and not f.startswith(("~$", ".")) and f not in SITE_SOURCES \
                    and re.match(r"\s*\d{3,}", f):
                names.append(stem)          # meter / Solar files start with a meter number
    return names


def _excel(src, dst, an):
    import openpyxl
    wb = openpyxl.load_workbook(src)
    for ws in wb.worksheets:
        ws.title = an.text(ws.title)[:31]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    c.value = an.text(c.value)
    wb.save(dst)


def _read_text(p):
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            with open(p, encoding=enc, newline="") as fh:
                return fh.read(), enc
        except UnicodeDecodeError:
            continue
    return None, None


def process(mode, src, dst, an):
    src = os.path.abspath(src)
    dst = os.path.abspath(dst)
    if os.path.commonpath([src, dst]) == src and os.path.isdir(src):
        sys.exit("DEST must not be inside SOURCE.")
    written, pictures, skipped, data_copies = [], [], [], []
    me = os.path.abspath(__file__)
    files = [(os.path.dirname(src), os.path.basename(src))] if os.path.isfile(src) else []
    if os.path.isdir(src):
        for dirpath, dirs, fs in os.walk(src):
            dirs[:] = sorted(d for d in dirs if d not in SKIP_DIRS)
            files += [(dirpath, f) for f in sorted(fs)]
    base = src if os.path.isdir(src) else os.path.dirname(src)
    for dirpath, f in files:
        if f.startswith("~$") or f == "anonymisation_key.csv" or os.path.abspath(os.path.join(dirpath, f)) == me:
            continue
        rel = os.path.relpath(os.path.join(dirpath, f), base)
        out = os.path.join(dst, an.text(rel))
        os.makedirs(os.path.dirname(out), exist_ok=True)
        p = os.path.join(dirpath, f)
        ext = os.path.splitext(f)[1].lower()
        if ext in TEXT_EXT:
            s, enc = _read_text(p)
            if s is None:
                skipped.append(rel)
                continue
            with open(out, "w", encoding=enc, newline="") as fh:
                fh.write(an.text(s))
            written.append(out)
        elif ext in EXCEL_EXT:
            if mode == "workspace" and (re.match(r"\s*\d{3,}", f) or os.path.getsize(p) > 5e6):
                shutil.copy2(p, out)          # meter / Solar / grid data: only the name is changed
                data_copies.append(out)
            else:
                _excel(p, out, an)
                written.append(out)
        elif ext in PICTURE_EXT:
            pictures.append(rel)
        elif mode == "workspace":
            shutil.copy2(p, out)
            data_copies.append(out)
        else:
            skipped.append(rel)
    return written, pictures, skipped, data_copies


def leftover_check(paths, an):
    found = []
    for p in paths:
        ext = os.path.splitext(p)[1].lower()
        name_hits = an.leftovers(os.path.basename(p))
        if name_hits:
            found.append((p, 0, name_hits))
        if ext in TEXT_EXT:
            s, _ = _read_text(p)
            for i, line in enumerate((s or "").splitlines(), 1):
                h = an.leftovers(line)
                if h:
                    found.append((p, i, h))
        elif ext in EXCEL_EXT:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(p, read_only=True)
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        for v in row:
                            if isinstance(v, str) and an.leftovers(v):
                                found.append((p, ws.title, an.leftovers(v)))
            except Exception:
                pass
    return found


def main(argv):
    import argparse
    ap = argparse.ArgumentParser(description="Anonymise site names (see the top of this file).")
    ap.add_argument("mode", choices=["results", "workspace"])
    ap.add_argument("source")
    ap.add_argument("dest")
    ap.add_argument("--sites", action="append", default=[],
                    help="extra site_comparison.csv (or folder with meter files) to learn the site names "
                         "from -- needed in 'results' mode when SOURCE has no site_comparison.csv")
    a = ap.parse_args(argv[1:])
    mode, src, dst = a.mode, a.source, a.dest
    if not os.path.exists(src):
        sys.exit(f"{src!r} does not exist")
    an = Anonymiser()
    names = collect_names(src)
    for extra in a.sites:
        names += collect_names(extra)
    an.learn(names)
    n_list = sorted({c for _, c, h in an.key_rows if h == "address in the list"})
    n_other = sorted(set(an.other.values()))
    print(f"[anonymise] sites found: {len(n_list)} from the list ({', '.join(n_list)}), "
          f"{len(n_other)} not in the list -> {OTHER_PREFIX} 1..{len(n_other)}")
    written, pictures, skipped, data_copies = process(mode, src, dst, an)
    keyp = os.path.join(src if os.path.isdir(src) else os.path.dirname(os.path.abspath(src)),
                        "anonymisation_key.csv")
    with open(keyp, "w", encoding="utf-8-sig") as fh:
        fh.write("real;code;how\n")
        for a, comp, c in SITES:
            fh.write(f"{a} ({comp});{code_of(c)};the list\n")
        for real, code, how in sorted(an.key_rows, key=lambda r: (r[1], r[0])):
            fh.write(f"{real};{code};{how}\n")
    print(f"[anonymise] {len(written)} file(s) anonymised into {dst}"
          + (f"; {len(data_copies)} data file(s) copied with an anonymised NAME only (meter, Solar, grid, "
             f"weather -- their contents are raw data: keep DEST private, share only the results made "
             f"in it)" if data_copies else ""))
    print(f"[anonymise] key (KEEP PRIVATE, not in DEST): {keyp}")
    if pictures:
        print(f"[anonymise] {len(pictures)} picture(s) NOT copied -- they may show real names; "
              f"make them again from the workspace copy (mode 'workspace'):")
        for p in pictures[:30]:
            print(f"    {p}")
        if len(pictures) > 30:
            print(f"    ... and {len(pictures) - 30} more")
    if skipped:
        print(f"[anonymise] {len(skipped)} file(s) NOT copied (type not handled): {skipped[:10]}")
    found = leftover_check(written, an)
    if found:
        print(f"\n[leftover check] {len(found)} place(s) still look like a real name or meter number:")
        for p, line, h in found[:50]:
            print(f"    {os.path.relpath(p, dst)} line {line}: {sorted(set(h))}")
        print("  Check them by hand before sharing (a long number can also be a timestamp or an id "
              "that is not a meter -- then it is fine).")
    else:
        print("[leftover check] no street name, company name or long meter number left in the files written.")
    return found


if __name__ == "__main__":
    main(sys.argv)
